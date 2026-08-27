from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from template_filler.services import markdown_to_excel, markdown_to_word


MARKDOWN = """# Quarterly Report

Revenue was **strong** and costs were *stable*.

- North region
- South region

> Reviewed by Finance.

| Region | Revenue |
| --- | ---: |
| North | 120 |
| South | 95 |
"""


def test_markdown_to_word_preserves_common_blocks():
    document = Document(BytesIO(markdown_to_word(MARKDOWN)))

    assert document.paragraphs[0].text == "Quarterly Report"
    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.paragraphs[1].text == "Revenue was strong and costs were stable."
    assert document.paragraphs[1].runs[1].bold is True
    assert document.paragraphs[1].runs[3].italic is True
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["Region", "Revenue"]
    assert [cell.text for cell in document.tables[0].rows[2].cells] == ["South", "95"]


def test_markdown_to_excel_creates_one_sheet_per_table():
    markdown = """# Revenue

| Region | Amount |
| --- | ---: |
| North | 120 |

# Costs

| Category | Amount |
| --- | ---: |
| Hosting | 30 |
"""
    workbook = load_workbook(BytesIO(markdown_to_excel(markdown)))

    assert workbook.sheetnames == ["Revenue", "Costs"]
    assert list(workbook["Revenue"].values) == [("Region", "Amount"), ("North", "120")]
    assert list(workbook["Costs"].values) == [("Category", "Amount"), ("Hosting", "30")]
    assert workbook["Revenue"]["A1"].font.bold is True
    assert workbook["Revenue"].freeze_panes == "A2"


def test_markdown_to_excel_exports_non_table_content():
    workbook = load_workbook(BytesIO(markdown_to_excel("# Notes\n\nA short paragraph.")))

    assert workbook.sheetnames == ["Content"]
    assert list(workbook.active.values) == [
        ("Type", "Content"),
        ("Heading 1", "Notes"),
        ("Paragraph", "A short paragraph."),
    ]


@pytest.mark.parametrize("converter", [markdown_to_word, markdown_to_excel])
def test_markdown_conversion_rejects_empty_content(converter):
    with pytest.raises(ValueError, match="must not be empty"):
        converter("  \n")
