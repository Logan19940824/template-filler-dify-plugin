from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook

from template_filler.services import fill_excel, fill_word
from template_filler.services.common import array_fields, expand_contexts, render_context


VALUES = {
    "modules": [
        {"name": "用户端", "features": [{"name": "登录"}, {"name": "注册"}]},
        {"name": "管理端", "features": [{"name": "审核"}]},
    ]
}

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _write_docx(path: Path, paragraph: str | list[str]) -> None:
    runs = paragraph if isinstance(paragraph, list) else [paragraph]
    run_xml = "".join(f"<w:r><w:t>{run}</w:t></w:r>" for run in runs)
    document_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}">
  <w:body><w:p>{run_xml}</w:p><w:sectPr/></w:body>
</w:document>'''
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document_xml)


def _docx_paragraphs(path: Path) -> list[str]:
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    return ["".join(node.text or "" for node in paragraph.iter(f"{{{W_NS}}}t")) for paragraph in root.iter(f"{{{W_NS}}}p")]


def test_context_and_indexes():
    text = "{{modules[]._index}}.{{modules[].name}}-{{modules[].features[]._index}}.{{modules[].features[].name}}"
    contexts = expand_contexts(VALUES, array_fields(text))
    assert [render_context(text, VALUES, context) for context in contexts] == [
        "1.用户端-1.登录",
        "1.用户端-2.注册",
        "2.管理端-1.审核",
    ]


def test_excel_nested_loop(tmp_path: Path):
    source, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["模块序号", "模块", "功能序号", "功能点"])
    sheet.append(["{{modules[]._index}}", "{{modules[].name}}", "{{modules[].features[]._index}}", "{{modules[].features[].name}}"])
    workbook.save(source)
    fill_excel(source, output, VALUES)
    rows = list(load_workbook(output).active.values)
    assert rows[1:] == [("1", "用户端", "1", "登录"), ("1", "用户端", "2", "注册"), ("2", "管理端", "1", "审核")]


def test_word_nested_loop(tmp_path: Path):
    source, output = tmp_path / "template.docx", tmp_path / "output.docx"
    _write_docx(source, "{{modules[]._index}}.{{modules[].name}}-{{modules[].features[]._index}}.{{modules[].features[].name}}")
    fill_word(source, output, VALUES)
    assert _docx_paragraphs(output) == [
        "1.用户端-1.登录",
        "1.用户端-2.注册",
        "2.管理端-1.审核",
    ]


def test_word_cross_run_replacement(tmp_path: Path):
    source, output = tmp_path / "template.docx", tmp_path / "output.docx"
    _write_docx(source, ["Order reference: {{order_", "id}}"])
    fill_word(source, output, {"order_id": "A-1"})
    assert _docx_paragraphs(output) == ["Order reference: A-1"]
