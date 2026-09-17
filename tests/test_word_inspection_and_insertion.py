from pathlib import Path
from copy import deepcopy
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from template_filler.services.word import fill_word, inspect_word_document, insert_word_placeholders, parse_word


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _document(path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Project: ")
    value = paragraph.add_run("          ")
    value.underline = True
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "No."
    table.cell(0, 1).text = "Name"
    document.save(path)


def _paragraph_runs(path: Path, paragraph_index: int) -> list[tuple[str, str | None]]:
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    paragraphs = root.findall(f".//{{{W_NS}}}body/{{{W_NS}}}p")
    result = []
    for run in paragraphs[paragraph_index].findall(f"./{{{W_NS}}}r"):
        text = "".join(node.text or "" for node in run.iter(f"{{{W_NS}}}t"))
        underline = run.find(f"./{{{W_NS}}}rPr/{{{W_NS}}}u")
        result.append((text, underline.get(f"{{{W_NS}}}val", "single") if underline is not None else None))
    return result


def test_inspect_and_insert_word_placeholders(tmp_path: Path):
    source, template, filled = tmp_path / "source.docx", tmp_path / "template.docx", tmp_path / "filled.docx"
    _document(source)
    inspected = inspect_word_document(source)

    assert inspected["paragraphs"][0]["text"] == "Project:           "
    assert inspected["paragraphs"][0]["blank_spans"][-1] == {"start": 8, "end": 19}
    assert len(inspected["tables"][0]["rows"]) == 3
    assert inspected["tables"][0]["rows"][1]["is_empty"] is True

    operations = {
        "document_sha256": inspected["document_sha256"],
        "operations": [
            {
                "type": "replace_text_range",
                "location": {"paragraph": 0},
                "expected_text": "Project:           ",
                "start": 9,
                "end": 19,
                "placeholder": "{{project}}",
            },
            {
                "type": "replace_text_range",
                "location": {"table": 0, "row": 1, "cell": 0, "paragraph": 0},
                "expected_text": "",
                "start": 0,
                "end": 0,
                "placeholder": "{{items[]._index}}",
            },
            {
                "type": "replace_text_range",
                "location": {"table": 0, "row": 1, "cell": 1, "paragraph": 0},
                "expected_text": "",
                "start": 0,
                "end": 0,
                "placeholder": "{{items[].name}}",
            },
            {"type": "delete_empty_rows", "table": 0, "rows": [2]},
        ],
    }
    insert_word_placeholders(source, template, operations)

    parsed = parse_word(template)
    assert parsed["values"] == {"project": "", "items": []}
    fill_word(template, filled, {"project": "Alpha", "items": [{"name": "A"}, {"name": "B"}]})
    document = Document(filled)
    assert document.paragraphs[0].text == "Project: Alpha"
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["No.", "Name"],
        ["1", "A"],
        ["2", "B"],
    ]
    assert _paragraph_runs(filled, 0)[-1] == ("Alpha", "single")


def test_insert_word_placeholders_rejects_stale_or_unsafe_operations(tmp_path: Path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    _document(source)
    inspected = inspect_word_document(source)

    with pytest.raises(ValueError, match="document_sha256"):
        insert_word_placeholders(source, output, {"document_sha256": "stale", "operations": [{}]})
    with pytest.raises(ValueError, match="expected text mismatch"):
        insert_word_placeholders(source, output, {
            "document_sha256": inspected["document_sha256"],
            "operations": [{
                "type": "replace_text_range",
                "location": {"paragraph": 0},
                "expected_text": "wrong",
                "start": 0,
                "end": 0,
                "placeholder": "{{x}}",
            }],
        })
    with pytest.raises(ValueError, match="non-empty or complex row"):
        insert_word_placeholders(source, output, {
            "document_sha256": inspected["document_sha256"],
            "operations": [{"type": "delete_empty_rows", "table": 0, "rows": [0]}],
        })


@pytest.mark.parametrize("count", [0, 1, 3, 12])
def test_word_loop_preserves_cross_run_formatting(tmp_path: Path, count: int):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Before ").italic = True
    paragraph.add_run("{{pro").underline = True
    paragraph.add_run("ject}}").bold = True
    paragraph.add_run(" after").italic = True
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "No.", "Name"
    table.cell(1, 0).text = "{{items[]._index}}"
    name = table.cell(1, 1).paragraphs[0]
    name.add_run("{{items[].").underline = True
    name.add_run("name}}").bold = True
    document.save(source)
    fill_word(source, output, {"project": "Alpha", "items": [{"name": f"Tenant {i}"} for i in range(count)]})
    result = Document(output)
    assert result.paragraphs[0].text == "Before Alpha after"
    assert [(run.text, run.italic, run.underline) for run in result.paragraphs[0].runs if run.text] == [
        ("Before ", True, None), ("Alpha", None, True), (" after", True, None),
    ]
    assert len(result.tables[0].rows) == count + 1
    for i, row in enumerate(result.tables[0].rows[1:]):
        assert [cell.text for cell in row.cells] == [str(i + 1), f"Tenant {i}"]
        assert row.cells[1].paragraphs[0].runs[0].underline is True


@pytest.mark.parametrize("change", [
    {"location": {"table": 0, "row": -1, "cell": 0, "paragraph": 0}},
    {"location": {"paragraph": True}},
    {"start": -1}, {"start": True}, {"end": 999},
    {"placeholder": "{{ }}"}, {"placeholder": "{{x}}{{y}}"},
    {"placeholder": "x{{x}}"}, {"placeholder": "{{x\ny}}"},
])
def test_invalid_word_range_never_writes_output(tmp_path: Path, change: dict):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    _document(source)
    inspection = inspect_word_document(source)
    operation = {
        "type": "replace_text_range", "location": {"paragraph": 0},
        "expected_text": inspection["paragraphs"][0]["text"],
        "start": 9, "end": 19, "placeholder": "{{project}}",
    }
    operation.update(change)
    output.write_bytes(b"existing output")
    with pytest.raises(ValueError):
        insert_word_placeholders(source, output, {"document_sha256": inspection["document_sha256"], "operations": [operation]})
    assert output.read_bytes() == b"existing output"


@pytest.mark.parametrize("kind", ["overlap", "same_start", "duplicate_delete", "replace_deleted"])
def test_conflicting_word_operations(tmp_path: Path, kind: str):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    _document(source)
    inspection = inspect_word_document(source)
    operation = {
        "type": "replace_text_range", "location": {"paragraph": 0},
        "expected_text": inspection["paragraphs"][0]["text"],
        "start": 9, "end": 19, "placeholder": "{{project}}",
    }
    second = deepcopy(operation)
    second.update(start=10 if kind == "overlap" else 9, end=10 if kind == "overlap" else 9)
    operations = [operation, second]
    if kind == "duplicate_delete":
        operations = [{"type": "delete_empty_rows", "table": 0, "rows": [1, 1]}]
    elif kind == "replace_deleted":
        operation.update(location={"table": 0, "row": 1, "cell": 0, "paragraph": 0}, expected_text="", start=0, end=0)
        operations = [operation, {"type": "delete_empty_rows", "table": 0, "rows": [1]}]
    with pytest.raises(ValueError):
        insert_word_placeholders(source, output, {"document_sha256": inspection["document_sha256"], "operations": operations})
    assert not output.exists()


@pytest.mark.parametrize("tag", ["pict", "drawing", "fldChar", "tab", "footnoteReference"])
def test_non_text_content_is_not_an_empty_row(tmp_path: Path, tag: str):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Header"
    table.cell(1, 0).paragraphs[0].add_run()._r.append(OxmlElement(f"w:{tag}"))
    document.save(source)
    inspection = inspect_word_document(source)
    assert inspection["unsupported_features"]
    assert not inspection["tables"][0]["rows"][1]["is_empty"]
    with pytest.raises(ValueError, match="non-empty or complex"):
        insert_word_placeholders(source, output, {
            "document_sha256": inspection["document_sha256"],
            "operations": [{"type": "delete_empty_rows", "table": 0, "rows": [1]}],
        })


def test_word_preserves_namespaces_parts_and_original_coordinates(tmp_path: Path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    _document(source)
    inspection = inspect_word_document(source)
    assert inspection["blocks"] == [{"type": "paragraph", "paragraph": 0}, {"type": "table", "table": 0}]
    assert len(inspection["tables"][0]["grid_widths_twips"]) == 2
    insert_word_placeholders(source, output, {
        "document_sha256": inspection["document_sha256"],
        "operations": [
            {"type": "delete_empty_rows", "table": 0, "rows": [1]},
            {"type": "replace_text_range", "location": {"table": 0, "row": 2, "cell": 0, "paragraph": 0},
             "expected_text": "", "start": 0, "end": 0, "placeholder": "{{name}}"},
        ],
    })
    assert Document(output).tables[0].cell(1, 0).text == "{{name}}"
    with ZipFile(source) as before, ZipFile(output) as after:
        assert set(before.namelist()) == set(after.namelist())
        for part in before.namelist():
            if part != "word/document.xml":
                assert before.read(part) == after.read(part)
        assert etree.fromstring(before.read("word/document.xml")).nsmap == etree.fromstring(after.read("word/document.xml")).nsmap


@pytest.mark.parametrize("count", [0, 1, 3, 12])
def test_tenant_sample_end_to_end(tmp_path: Path, count: int):
    source = Path(__file__).resolve().parents[1] / "docs/目标-统一租户台账-Word版2.docx"
    if not source.exists():
        pytest.skip("User-supplied acceptance document is not part of the repository")
    inspection = inspect_word_document(source)
    assert len(inspection["paragraphs"]) == 6
    assert len(inspection["tables"][0]["rows"]) == 11
    assert len(inspection["tables"][0]["rows"][1]["cells"]) == 9
    assert not inspection["unsupported_features"]
    operations = []
    scalar_fields = {1: [(3, 3, "项目"), (32, 32, "制表日期")],
                     2: [(5, 5, "报送单位"), (30, 30, "联系人")],
                     5: [(4, 17, "制表人"), (22, 36, "审核"), (41, 52, "日期")]}
    values = {}
    for paragraph_index, fields in scalar_fields.items():
        paragraph = inspection["paragraphs"][paragraph_index]
        for start, end, field in fields:
            operations.append({"type": "replace_text_range", "location": paragraph["location"],
                               "expected_text": paragraph["text"], "start": start, "end": end,
                               "placeholder": "{{" + field + "}}"})
            values[field] = "示例日期" if "日期" in field else "示例" + field
    headers = [cell["paragraphs"][0]["text"] for cell in inspection["tables"][0]["rows"][0]["cells"]]
    for column, field in enumerate(headers):
        operations.append({"type": "replace_text_range",
                           "location": {"table": 0, "row": 1, "cell": column, "paragraph": 0},
                           "expected_text": "", "start": 0, "end": 0,
                           "placeholder": "{{租户[]." + ("_index" if column == 0 else field) + "}}"})
    operations.append({"type": "delete_empty_rows", "table": 0, "rows": list(range(2, 11))})
    template, output = tmp_path / "template.docx", tmp_path / "output.docx"
    insert_word_placeholders(source, template, {"document_sha256": inspection["document_sha256"], "operations": operations})
    assert len(parse_word(template)["placeholders"]) == 16
    values["租户"] = [{field: f"测试{i + 1}{field}" for field in headers[1:]} for i in range(count)]
    fill_word(template, output, values)
    document = Document(output)
    assert len(document.tables[0].rows) == count + 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == headers
    assert not parse_word(output)["placeholders"]
    for index, row in enumerate(document.tables[0].rows[1:], 1):
        assert row.cells[0].text == str(index)
        assert row.cells[2].text == f"测试{index}公司名"
    for field in ("制表人", "审核"):
        run = next(run for run in document.paragraphs[5].runs if run.text == values[field])
        assert run.underline is True
    with ZipFile(source) as before, ZipFile(output) as after:
        original = etree.fromstring(before.read("word/document.xml"))
        result = etree.fromstring(after.read("word/document.xml"))
        for xpath in (f".//{{{W_NS}}}tblPr", f".//{{{W_NS}}}tblGrid", f".//{{{W_NS}}}sectPr"):
            assert etree.tostring(original.find(xpath)) == etree.tostring(result.find(xpath))
