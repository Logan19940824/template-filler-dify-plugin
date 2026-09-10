from copy import copy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from template_filler.services.excel import insert_excel_placeholders


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quote"
    sheet["A1"] = "Customer"
    sheet["B1"] = "Example Co"
    sheet["B1"].font = copy(sheet["A1"].font)
    sheet["B2"] = "=SUM(1,2)"
    workbook.save(path)


def test_insert_excel_placeholders_changes_values_and_preserves_style(tmp_path: Path):
    source, output = tmp_path / "source.xlsx", tmp_path / "output.xlsx"
    _workbook(source)
    before = load_workbook(source).active["B1"]

    insert_excel_placeholders(
        source,
        output,
        [{"sheet": "Quote", "cell": "B1", "expected_value": "Example Co", "placeholder": "{{customer_name}}"}],
    )

    after = load_workbook(output).active["B1"]
    assert after.value == "{{customer_name}}"
    assert after.style_id == before.style_id


def test_insert_excel_placeholders_rejects_mismatch_and_formula(tmp_path: Path):
    source, output = tmp_path / "source.xlsx", tmp_path / "output.xlsx"
    _workbook(source)
    with pytest.raises(ValueError, match="expected value mismatch"):
        insert_excel_placeholders(source, output, [{"sheet": "Quote", "cell": "B1", "expected_value": "wrong", "placeholder": "{{x}}"}])
    with pytest.raises(ValueError, match="formula cell"):
        insert_excel_placeholders(source, output, [{"sheet": "Quote", "cell": "B2", "expected_value": "=SUM(1,2)", "placeholder": "{{total}}"}])


def test_insert_excel_placeholders_inserts_styled_row(tmp_path: Path):
    source, output = tmp_path / "source.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quote"
    sheet.append(["No.", "Product", "Total"])
    sheet.append(["Example", "Sample product", "=LEN(B2)"])
    sheet["B2"].font = copy(sheet["A1"].font)
    sheet.row_dimensions[2].height = 30
    sheet.append(["Subtotal", None, "=SUM(C2:C2)"])
    workbook.save(source)

    insert_excel_placeholders(
        source,
        output,
        [{
            "type": "insert_row",
            "sheet": "Quote",
            "before_row": 3,
            "expected_before": {"cell": "A3", "expected_value": "Subtotal"},
            "copy_style_from_row": 2,
            "cells": {"A": "{{items[]._index}}", "B": "{{items[].name}}"},
        }],
    )

    sheet = load_workbook(output).active
    assert [cell.value for cell in sheet[3]] == ["{{items[]._index}}", "{{items[].name}}", "=LEN(B3)"]
    assert sheet["B3"].style_id == sheet["B2"].style_id
    assert sheet.row_dimensions[3].height == 30
    assert sheet["A4"].value == "Subtotal"
