from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from markdown_it import MarkdownIt
from markdown_it.token import Token
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_MARKDOWN = MarkdownIt("commonmark", {"html": False}).enable("table")
_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


@dataclass
class MarkdownTable:
    title: str
    headers: list[str]
    rows: list[list[str]]


def _parse(markdown_content: str) -> list[Token]:
    if not isinstance(markdown_content, str) or not markdown_content.strip():
        raise ValueError("markdown_content must not be empty")
    return _MARKDOWN.parse(markdown_content)


def _inline_text(tokens: list[Token] | None) -> str:
    parts: list[str] = []
    for token in tokens or []:
        if token.type in {"text", "code_inline", "html_inline"}:
            parts.append(token.content)
        elif token.type in {"softbreak", "hardbreak"}:
            parts.append("\n")
        elif token.type == "image":
            label = token.content or "image"
            source = token.attrGet("src")
            parts.append(f"{label} ({source})" if source else label)
        elif token.children:
            parts.append(_inline_text(token.children))
    return "".join(parts).strip()


def _read_table(tokens: list[Token], start: int, title: str) -> tuple[MarkdownTable, int]:
    headers: list[str] = []
    rows: list[list[str]] = []
    row: list[str] = []
    cell = ""
    in_header = False
    index = start + 1
    while index < len(tokens):
        token = tokens[index]
        if token.type == "table_close":
            break
        if token.type == "thead_open":
            in_header = True
        elif token.type == "thead_close":
            in_header = False
        elif token.type == "tr_open":
            row = []
        elif token.type in {"th_open", "td_open"}:
            cell = ""
        elif token.type == "inline":
            cell = _inline_text(token.children)
        elif token.type in {"th_close", "td_close"}:
            row.append(cell)
        elif token.type == "tr_close":
            if in_header and not headers:
                headers = row
            else:
                rows.append(row)
        index += 1

    width = max([len(headers), *(len(row) for row in rows)], default=0)
    if len(headers) < width:
        headers.extend(f"Column {column}" for column in range(len(headers) + 1, width + 1))
    rows = [row + [""] * (width - len(row)) for row in rows]
    return MarkdownTable(title=title, headers=headers, rows=rows), index + 1


def _extract_tables(tokens: list[Token]) -> list[MarkdownTable]:
    tables: list[MarkdownTable] = []
    heading = ""
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "heading_open" and index + 1 < len(tokens):
            heading = _inline_text(tokens[index + 1].children)
        elif token.type == "table_open":
            table, index = _read_table(tokens, index, heading or f"Table {len(tables) + 1}")
            tables.append(table)
            continue
        index += 1
    return tables


def _apply_run_format(run, *, bold: bool, italic: bool, strike: bool, link: bool, code: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.strike = strike
    if code:
        run.font.name = "Courier New"
    if link:
        run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        run.underline = True


def _render_inline(paragraph, tokens: list[Token] | None) -> None:
    bold = italic = strike = False
    link = ""
    for token in tokens or []:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type == "s_open":
            strike = True
        elif token.type == "s_close":
            strike = False
        elif token.type == "link_open":
            link = token.attrGet("href") or ""
        elif token.type == "link_close":
            link = ""
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        elif token.type == "image":
            label = token.content or "image"
            source = token.attrGet("src")
            run = paragraph.add_run(f"{label} ({source})" if source else label)
            _apply_run_format(run, bold=bold, italic=italic, strike=strike, link=bool(source))
        elif token.type in {"text", "html_inline", "code_inline"}:
            run = paragraph.add_run(token.content)
            _apply_run_format(
                run,
                bold=bold,
                italic=italic,
                strike=strike,
                link=bool(link),
                code=token.type == "code_inline",
            )
        elif token.children:
            _render_inline(paragraph, token.children)


def _list_style(list_type: str, depth: int) -> str:
    base = "List Number" if list_type == "ordered" else "List Bullet"
    return base if depth == 1 else f"{base} {min(depth, 3)}"


def _add_horizontal_rule(document) -> None:
    paragraph = document.add_paragraph()
    paragraph_property = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B7C0CC")
    borders.append(bottom)
    paragraph_property.append(borders)


def markdown_to_word(markdown_content: str) -> bytes:
    tokens = _parse(markdown_content)
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10.5)

    list_stack: list[str] = []
    quote_depth = 0
    pending_block: tuple[str, int] | None = None
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "table_open":
            markdown_table, index = _read_table(tokens, index, "")
            if not markdown_table.headers:
                continue
            table = document.add_table(rows=1, cols=len(markdown_table.headers))
            table.style = "Table Grid"
            for column, value in enumerate(markdown_table.headers):
                paragraph = table.rows[0].cells[column].paragraphs[0]
                run = paragraph.add_run(value)
                run.bold = True
            for values in markdown_table.rows:
                cells = table.add_row().cells
                for column, value in enumerate(values):
                    cells[column].text = value
            continue
        if token.type == "heading_open":
            pending_block = ("heading", int(token.tag[1:]))
        elif token.type == "paragraph_open":
            pending_block = ("paragraph", 0)
        elif token.type == "inline" and pending_block:
            kind, level = pending_block
            if kind == "heading":
                paragraph = document.add_heading("", level=level)
            elif list_stack:
                paragraph = document.add_paragraph(style=_list_style(list_stack[-1], len(list_stack)))
            elif quote_depth:
                paragraph = document.add_paragraph(style="Intense Quote")
            else:
                paragraph = document.add_paragraph()
            _render_inline(paragraph, token.children)
            pending_block = None
        elif token.type == "bullet_list_open":
            list_stack.append("bullet")
        elif token.type == "ordered_list_open":
            list_stack.append("ordered")
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            list_stack.pop()
        elif token.type == "blockquote_open":
            quote_depth += 1
        elif token.type == "blockquote_close":
            quote_depth -= 1
        elif token.type in {"fence", "code_block"}:
            paragraph = document.add_paragraph(style="No Spacing")
            run = paragraph.add_run(token.content.rstrip("\n"))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif token.type == "hr":
            _add_horizontal_rule(document)
        elif token.type == "html_block" and token.content.strip():
            document.add_paragraph(token.content.strip())
        index += 1

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _outline_rows(tokens: list[Token]) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    lists: list[str] = []
    quoted = False
    pending = ""
    for token in tokens:
        if token.type == "heading_open":
            pending = f"Heading {token.tag[1:]}"
        elif token.type == "paragraph_open":
            pending = "List item" if lists else "Quote" if quoted else "Paragraph"
        elif token.type == "inline" and pending:
            rows.append((pending, _inline_text(token.children)))
            pending = ""
        elif token.type == "bullet_list_open":
            lists.append("bullet")
        elif token.type == "ordered_list_open":
            lists.append("ordered")
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif token.type == "blockquote_open":
            quoted = True
        elif token.type == "blockquote_close":
            quoted = False
        elif token.type in {"fence", "code_block"}:
            rows.append(("Code", token.content.rstrip("\n")))
    return rows


def _sheet_name(value: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("_", value).strip(" '")[:31] or "Sheet"
    candidate = base
    number = 2
    while candidate.casefold() in used:
        suffix = f" {number}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        number += 1
    used.add(candidate.casefold())
    return candidate


def _format_worksheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        width = max((len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)), default=0)
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 60)


def markdown_to_excel(markdown_content: str) -> bytes:
    tokens = _parse(markdown_content)
    tables = _extract_tables(tokens)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()

    if tables:
        for number, table in enumerate(tables, 1):
            sheet = workbook.create_sheet(_sheet_name(table.title or f"Table {number}", used_names))
            sheet.append(table.headers)
            for row in table.rows:
                sheet.append(row)
            _format_worksheet(sheet)
    else:
        sheet = workbook.create_sheet("Content")
        sheet.append(["Type", "Content"])
        for row in _outline_rows(tokens):
            sheet.append(row)
        _format_worksheet(sheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
