from __future__ import annotations

from copy import copy
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string

from .common import array_fields, expand_contexts, placeholder_names, render_context, replace_text, value_defaults


def _open_excel(path: Path, *, read_only: bool = False):
    try:
        return load_workbook(path, read_only=read_only, data_only=False)
    except Exception as exc:
        raise ValueError("template is not a valid Excel XLSX file or is corrupted") from exc


def parse_excel(path: Path) -> dict[str, Any]:
    workbook = _open_excel(path, read_only=True)
    found: dict[str, list[dict[str, Any]]] = {}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for name in placeholder_names(cell.value):
                        found.setdefault(name, []).append({"sheet": sheet.title, "cells": [cell.coordinate]})
    return {"file_type": "xlsx", "placeholders": [{"name": n, "occurrences": o} for n, o in found.items()], "values": value_defaults(list(found))}


def _json_value(value: Any) -> Any:
    return value.isoformat() if isinstance(value, (date, datetime, time)) else value


def inspect_excel_workbook(path: Path) -> dict[str, Any]:
    workbook = _open_excel(path)
    sheets = []
    for sheet in workbook.worksheets:
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                item = {
                    "coordinate": cell.coordinate,
                    "value": _json_value(cell.value),
                    "data_type": cell.data_type,
                    "style_id": cell.style_id,
                }
                if cell.data_type == "f":
                    item["formula"] = cell.value
                cells.append(item)
        sheets.append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                "merged_cells": [str(range_) for range_ in sheet.merged_cells.ranges],
                "cells": cells,
            }
        )
    return {"file_type": "xlsx", "sheets": sheets}


def _row_intersects_merge(sheet, row: int) -> bool:
    return any(merged.min_row <= row <= merged.max_row for merged in sheet.merged_cells.ranges)


def _replace_cell(sheet, operation: dict[str, Any]) -> None:
    coordinate = operation.get("cell")
    expected = operation.get("expected_value")
    placeholder = operation.get("placeholder")
    if not all(isinstance(value, str) and value.strip() for value in (coordinate, placeholder)):
        raise ValueError("replace_cell requires cell and placeholder strings")
    if not placeholder.startswith("{{") or not placeholder.endswith("}}") or not placeholder[2:-2].strip():
        raise ValueError(f"invalid placeholder: {placeholder}")
    try:
        cell = sheet[coordinate]
    except (KeyError, ValueError, TypeError) as exc:
        raise ValueError(f"invalid cell coordinate: {sheet.title}!{coordinate}") from exc
    if cell.data_type in {"f", "e"}:
        cell_type = "formula" if cell.data_type == "f" else "error"
        raise ValueError(f"cannot replace {cell_type} cell: {sheet.title}!{coordinate}")
    for merged_range in sheet.merged_cells.ranges:
        if cell.coordinate in merged_range and cell.coordinate != merged_range.start_cell.coordinate:
            raise ValueError(f"cannot replace non-top-left merged cell: {sheet.title}!{coordinate}")
    if cell.value != expected:
        raise ValueError(
            f"expected value mismatch at {sheet.title}!{coordinate}: "
            f"expected {expected!r}, found {cell.value!r}"
        )
    cell.value = placeholder


def _insert_row(sheet, operation: dict[str, Any]) -> None:
    before_row = operation.get("before_row")
    source_row = operation.get("copy_style_from_row")
    anchor = operation.get("expected_before")
    cells = operation.get("cells")
    if not isinstance(before_row, int) or not isinstance(source_row, int) or not isinstance(anchor, dict) or not isinstance(cells, dict):
        raise ValueError("insert_row requires before_row, copy_style_from_row, expected_before, and cells")
    if not 1 <= before_row <= sheet.max_row + 1 or not 1 <= source_row <= sheet.max_row:
        raise ValueError(f"invalid insert_row position in worksheet: {sheet.title}")
    anchor_cell, anchor_value = anchor.get("cell"), anchor.get("expected_value")
    if not isinstance(anchor_cell, str) or not anchor_cell.strip():
        raise ValueError("insert_row expected_before requires a cell")
    try:
        if sheet[anchor_cell].row != before_row:
            raise ValueError("insert_row expected_before cell must be on before_row")
    except (KeyError, ValueError, TypeError) as exc:
        raise ValueError(f"invalid insert_row anchor: {sheet.title}!{anchor_cell}") from exc
    if sheet[anchor_cell].value != anchor_value:
        raise ValueError(f"insert_row anchor mismatch at {sheet.title}!{anchor_cell}")
    if _row_intersects_merge(sheet, source_row) or (before_row <= sheet.max_row and _row_intersects_merge(sheet, before_row)):
        raise ValueError(f"insert_row cannot use merged rows in worksheet: {sheet.title}")

    placeholders: dict[int, str] = {}
    for column, placeholder in cells.items():
        if not isinstance(column, str) or not isinstance(placeholder, str) or not placeholder.startswith("{{") or not placeholder.endswith("}}") or not placeholder[2:-2].strip():
            raise ValueError("insert_row cells must map column letters to non-empty placeholders")
        try:
            placeholders[column_index_from_string(column.upper())] = placeholder
        except ValueError as exc:
            raise ValueError(f"invalid insert_row column: {column}") from exc

    snapshots = [(cell.value, copy(cell._style), cell.coordinate) for cell in sheet[source_row]]
    source_dimension = sheet.row_dimensions[source_row]
    sheet.insert_rows(before_row)
    target_dimension = sheet.row_dimensions[before_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed
    for column, (value, style, origin) in enumerate(snapshots, 1):
        target = sheet.cell(before_row, column)
        target._style = style
        if column in placeholders:
            target.value = placeholders[column]
        elif isinstance(value, str) and value.startswith("="):
            target.value = Translator(value, origin=origin).translate_formula(target.coordinate)


def insert_excel_placeholders(path: Path, output: Path, operations: list[dict[str, Any]]) -> None:
    if not isinstance(operations, list) or not operations:
        raise ValueError("operations must be a non-empty JSON array")
    workbook = _open_excel(path)
    seen: set[tuple[str, str]] = set()
    insert_sheets: set[str] = set()
    replacements = []
    insertions = []
    for operation in operations:
        if not isinstance(operation, dict):
            raise ValueError("each operation must be a JSON object")
        sheet_name = operation.get("sheet")
        operation_type = operation.get("type", "replace_cell")
        if not isinstance(sheet_name, str) or not sheet_name.strip():
            raise ValueError("each operation requires a sheet string")
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"worksheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
        if operation_type == "replace_cell":
            coordinate = operation.get("cell")
            if not isinstance(coordinate, str):
                raise ValueError("replace_cell requires a cell string")
            key = (sheet_name, coordinate.upper())
            if key in seen:
                raise ValueError(f"duplicate operation for {sheet_name}!{coordinate}")
            seen.add(key)
            replacements.append((sheet, operation))
        elif operation_type == "insert_row":
            if sheet_name in insert_sheets:
                raise ValueError(f"only one insert_row operation is allowed per worksheet: {sheet_name}")
            insert_sheets.add(sheet_name)
            insertions.append((sheet, operation))
        else:
            raise ValueError(f"unsupported operation type: {operation_type}")
    for sheet, operation in replacements:
        _replace_cell(sheet, operation)
    for sheet, operation in insertions:
        _insert_row(sheet, operation)
    workbook.save(output)


def _copy_row_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.fill = copy(source.fill)
    target.font = copy(source.font)
    target.border = copy(source.border)
    target.protection = copy(source.protection)


def fill_excel(path: Path, output: Path, values: dict[str, Any]) -> None:
    workbook = _open_excel(path)
    for sheet in workbook.worksheets:
        for row_index in range(sheet.max_row, 0, -1):
            row = list(sheet[row_index])
            fields = list(dict.fromkeys(f for cell in row if isinstance(cell.value, str) for f in array_fields(cell.value)))
            if fields:
                contexts = expand_contexts(values, fields)
                snapshots = [(cell.value, cell) for cell in row]
                if not contexts:
                    sheet.delete_rows(row_index)
                    continue
                for offset, context in enumerate(contexts):
                    target_index = row_index + offset
                    if offset:
                        sheet.insert_rows(target_index)
                    for column, (value, source) in enumerate(snapshots, 1):
                        target = sheet.cell(target_index, column)
                        target.value = render_context(value, values, context) if isinstance(value, str) else value
                        _copy_row_style(source, target)
            else:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = replace_text(cell.value, values)
    workbook.save(output)
