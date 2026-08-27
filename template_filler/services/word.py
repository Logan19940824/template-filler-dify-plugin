from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document

from .common import array_fields, expand_contexts, placeholder_names, render_context, replace_text, value_defaults


def _open_word(path: Path) -> Document:
    try:
        return Document(path)
    except Exception as exc:
        raise ValueError("template is not a valid Word DOCX file or is corrupted") from exc


def _paragraph_text(paragraph) -> str:
    return "".join(run.text or "" for run in paragraph.runs)


def _set_xml_text(element, text: str) -> None:
    nodes = [node for node in element.iter() if node.tag.endswith("}t")]
    if nodes:
        nodes[0].text = text
        for node in nodes[1:]:
            node.text = ""


def _replace_paragraph(paragraph, values: dict[str, Any]) -> None:
    text = _paragraph_text(paragraph)
    replaced = replace_text(text, values)
    if replaced == text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replaced)


def _iter_paragraphs(document):
    for index, paragraph in enumerate(document.paragraphs):
        yield paragraph, {"paragraph": index}
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    yield paragraph, {"table": table_index, "row": row_index, "cell": cell_index, "paragraph": paragraph_index}


def parse_word(path: Path) -> dict[str, Any]:
    document = _open_word(path)
    found: dict[str, list[dict[str, Any]]] = {}
    for paragraph, location in _iter_paragraphs(document):
        for name in placeholder_names(_paragraph_text(paragraph)):
            found.setdefault(name, []).append(location)
    return {"file_type": "docx", "placeholders": [{"name": n, "occurrences": o} for n, o in found.items()], "values": value_defaults(list(found))}


def fill_word(path: Path, output: Path, values: dict[str, Any]) -> None:
    document = _open_word(path)
    for paragraph in list(document.paragraphs):
        fields = array_fields(_paragraph_text(paragraph))
        if not fields:
            _replace_paragraph(paragraph, values)
            continue
        contexts = expand_contexts(values, fields)
        anchor = paragraph._p
        template_element = deepcopy(paragraph._p)
        if not contexts:
            paragraph._p.getparent().remove(paragraph._p)
            continue
        for index, context in enumerate(contexts):
            target = paragraph._p if index == 0 else deepcopy(template_element)
            if index:
                anchor.addnext(target)
                anchor = target
            text = _paragraph_text(paragraph) if index == 0 else "".join(node.text or "" for node in target.iter() if node.tag.endswith("}t"))
            _set_xml_text(target, render_context(text, values, context))
    for table in document.tables:
        for row in list(table.rows):
            paragraphs = [p for cell in row.cells for p in cell.paragraphs]
            fields = list(dict.fromkeys(f for p in paragraphs for f in array_fields(_paragraph_text(p))))
            if not fields:
                for paragraph in paragraphs:
                    _replace_paragraph(paragraph, values)
                continue
            contexts = expand_contexts(values, fields)
            anchor = row._tr
            template_row = deepcopy(row._tr)
            if not contexts:
                row._tr.getparent().remove(row._tr)
                continue
            for index, context in enumerate(contexts):
                target = row._tr if index == 0 else deepcopy(template_row)
                if index:
                    anchor.addnext(target)
                    anchor = target
                for node in target.iter():
                    if node.tag.endswith("}t"):
                        node.text = render_context(node.text or "", values, context)
    document.save(output)
