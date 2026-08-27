from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


ROOT = Path(__file__).parent


def build_docx():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Template Filler Word Test")
    run.bold = True
    run.font.size = Pt(18)
    document.add_paragraph("Customer: {{customer_name}}")
    # Deliberately split across runs to test Word's cross-run replacement.
    split = document.add_paragraph()
    split.add_run("Order reference: {{order_")
    split.add_run("id}}")
    document.add_paragraph("Modules and features:")
    document.add_paragraph("{{modules[]._index}}. {{modules[].name}} / {{modules[].features[]._index}}. {{modules[].features[].name}} / price={{modules[].features[].price}}")

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Product", "Quantity", "Price"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
    row = table.add_row()
    row.cells[0].text = "{{modules[].name}}"
    row.cells[1].text = "{{modules[].features[]._index}}"
    row.cells[2].text = "{{modules[].features[].name}}"
    document.add_paragraph("Missing value remains: {{not_provided}}")
    document.save(ROOT / "template.docx")


def build_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(bottom=Side(style="thin", color="B7C9D6"))
    sheet["A1"] = "Excel Template Filler Test"
    sheet["A2"] = "Customer"
    sheet["B2"] = "{{customer_name}}"
    sheet["A3"] = "Order ID"
    sheet["B3"] = "{{order_id}}"
    for cell in sheet[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    for cell, value in zip(sheet[4], ["Product", "Quantity", "Price"]):
        cell.value = value
    sheet["A5"] = "{{modules[]._index}}"
    sheet["B5"] = "{{modules[].name}}"
    sheet["C5"] = "{{modules[].features[]._index}}"
    sheet["D5"] = "{{modules[].features[].name}}"
    sheet["E5"] = "{{modules[].features[].price}}"
    sheet["A6"] = "Missing: {{not_provided}}"
    for row in sheet.iter_rows(min_row=1, max_row=6, min_col=1, max_col=3):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 28
    sheet.column_dimensions["E"].width = 16

    detail = workbook.create_sheet("Details")
    detail["A1"] = "Sheet-specific placeholder"
    detail["B1"] = "{{shipping_address}}"
    detail["A3"] = "{{modules[]._index}}"
    detail["B3"] = "{{modules[].name}}"
    detail["C3"] = "{{modules[].features[]._index}}"
    detail["D3"] = "{{modules[].features[].name}}"
    detail.column_dimensions["A"].width = 28
    detail.column_dimensions["B"].width = 28
    detail.column_dimensions["C"].width = 16
    detail.column_dimensions["D"].width = 28
    workbook.save(ROOT / "template.xlsx")


if __name__ == "__main__":
    build_docx()
    build_xlsx()
    print("created", ROOT / "template.docx", ROOT / "template.xlsx")
